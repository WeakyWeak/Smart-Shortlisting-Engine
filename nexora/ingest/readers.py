"""Stage 1 - pdf/docx/xml/txt to clean text. Every reader returns plain text so
Stage 2 treats all formats identically."""
from __future__ import annotations

import html
import re
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path

from .. import config

SUPPORTED = {".pdf", ".docx", ".txt", ".xml"}
READABLE = SUPPORTED | {".md"}   # .md is only ever a JD, never a resume

# Ligatures and smart quotes survive PDF extraction and break token matching.
_TYPO_FIXES = {
    "ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl",
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "−": "-", " ": " ", "​": "",
    "•": "\n- ", "●": "\n- ", "▪": "\n- ", "‣": "\n- ", "·": "\n- ",
}


def clean(text: str) -> str:
    """Normalise text from any source, so a mangled PDF ends up matching the same
    resume as clean .txt."""
    if not text:
        return ""
    try:
        import ftfy
        text = ftfy.fix_text(text)
    except ImportError:
        pass
    text = html.unescape(text)
    for bad, good in _TYPO_FIXES.items():
        text = text.replace(bad, good)
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.strip() for line in text.splitlines()).strip()
    # PDFs often put the bullet glyph on its own line and the text on the next.
    # Re-join, or every bullet parses as an empty marker plus an orphan sentence.
    text = re.sub(r"^[-*\u2022\u25cf\u25aa\u2023]\s*\n+(?=\S)", "- ", text, flags=re.M)
    return text.strip()


def _read_pdf(path: Path) -> str:
    import pymupdf
    with pymupdf.open(path) as doc:
        text = "\n".join(page.get_text() for page in doc)
    # Almost no text usually means a scan -- try OCR.
    if len(text.strip()) < config.OCR_MIN_CHARS:
        text = _ocr_pdf(path) or text
    return text


def _ocr_pdf(path: Path) -> str:
    """Fallback for scanned resumes. Optional dep -- not worth blocking on."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    out = []
    try:
        with pymupdf.open(path) as doc:
            for page in doc:
                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                out.append(pytesseract.image_to_string(img))
    except Exception:
        return ""
    return "\n".join(out)


def _read_docx(path: Path) -> str:
    """Paragraphs plus tables -- skills are often in a table, and python-docx
    doesn't surface those in `.paragraphs`."""
    import docx
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _read_xml(path: Path) -> str:
    """Tag names carry section meaning, so emit them as headers. Gives Stage 2
    ground-truth sections, which validate_parser.py checks against."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return re.sub(r"<[^>]+>", " ", raw)

    lines: list[str] = []

    def walk(node: ET.Element, depth: int = 0) -> None:
        tag = node.tag.lower()
        text = (node.text or "").strip()
        children = list(node)
        if depth == 1 and children:
            # Second-level tags are the sections.
            lines.append(f"\n{_tag_to_header(tag)}")
        if text:
            lines.append("- " + text if tag == "bullet" else text)
        for key, val in node.attrib.items():
            if key == "category" and val:
                lines.append(f"{val}:")
        for child in children:
            walk(child, depth + 1)

    walk(root)
    return "\n".join(lines)


_TAG_HEADERS = {
    "applicant": "CONTACT", "summary": "PROFESSIONAL SUMMARY",
    "education": "EDUCATION", "technicalskills": "TECHNICAL SKILLS",
    "skills": "TECHNICAL SKILLS", "experience": "EXPERIENCE",
    "projects": "PROJECTS", "certifications": "CERTIFICATIONS",
    "achievements": "ACHIEVEMENTS", "activities": "ACTIVITIES",
}


def _tag_to_header(tag: str) -> str:
    return _TAG_HEADERS.get(tag, tag.upper())


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


_READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".xml": _read_xml,
            ".txt": _read_text, ".md": _read_text}


def read(path: Path) -> str:
    """Read any supported file. Returns '' rather than raising -- one bad file
    shouldn't kill the whole run."""
    suffix = path.suffix.lower()
    if suffix not in _READERS:
        return ""
    try:
        return clean(_READERS[suffix](path))
    except Exception as exc:  # noqa: BLE001 - broad on purpose, see docstring
        print(f"  ! failed to read {path.name}: {type(exc).__name__}: {exc}")
        return ""
